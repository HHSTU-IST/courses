"""把作业得分回填到成绩表。

从作业目录的已批改文件名（形如 ``01-2401143001-陈炯轩-60.md``）解析学号与得分，
按学号在成绩表中定位行，把得分写入指定列（默认「作业1」）。

用法（命令行参数保持纯 ASCII，中文路径在 cmd 层会被 GBK 编码截断）：
    micromamba run -n kaggle python scripts/fill_scores.py --class 241
    micromamba run -n kaggle python scripts/fill_scores.py --class 241 --task 2
    micromamba run -n kaggle python scripts/fill_scores.py --class 241 --dry-run

``--class 241`` 会自动定位 ``智能B241-*`` 目录与 ``智能B241`` 工作表；
也可直接给出目录位置参数（在 IDE 里跑时更方便）。

默认先把原表另存为 ``<原表>.bak.xlsx`` 再原地写入；
``--dry-run`` 只打印对应关系，不改动工作簿。

也可作为模块使用::

    from fill_scores import collect_scores, fill_column

    scores, skipped = collect_scores(Path("智能B242-01"))
    result = fill_column(
        "2026-cv-成绩.xlsx", "智能B242", "作业1", scores
    )
"""

from __future__ import annotations

import argparse
import shutil
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

DEFAULT_XLSX = "2026-cv-成绩.xlsx"


@dataclass
class FillResult:
    """一次回填的结果。"""

    sheet: str
    column: str
    col_index: int
    filled: int = 0
    backup: Path | None = None
    missing_in_sheet: list[int] = field(default_factory=list)
    """目录里有、成绩表里查不到学号的。"""
    missing_in_dir: list[tuple[int, str]] = field(default_factory=list)
    """成绩表里有、目录里没交的，元素为 (学号, 姓名)。"""


def resolve_dir(directory: str | Path | None, klass: str | None) -> Path:
    """定位作业目录：显式目录优先，否则按班级号 glob ``智能B{klass}-*``。"""
    if directory:
        return Path(directory).resolve()
    if klass:
        found = sorted(p for p in Path.cwd().glob(f"智能B{klass}-*") if p.is_dir())
        if len(found) != 1:
            raise SystemExit(f"班级 {klass} 匹配到 {len(found)} 个目录，请给出位置参数")
        return found[0].resolve()
    raise SystemExit("请给出作业目录，或使用 --class 指定班级号")


def parse_filename(stem: str) -> tuple[int, str, int] | None:
    """从 ``01-2401143001-陈炯轩-60`` 解析出 (学号, 姓名, 得分)。"""
    parts = stem.split("-")
    if len(parts) < 4:
        return None
    sid, score = parts[1], parts[-1]
    if not (sid.isdigit() and score.isdigit()):
        return None
    return int(sid), "-".join(parts[2:-1]), int(score)


def collect_scores(directory: Path) -> tuple[dict[int, tuple[str, int]], list[str]]:
    """收集目录下所有 md 文件的得分，返回 ({学号: (姓名, 得分)}, 无法解析的文件名)。"""
    scores: dict[int, tuple[str, int]] = {}
    skipped: list[str] = []
    for path in sorted(p for p in directory.glob("*.md") if p.is_file()):
        parsed = parse_filename(path.stem)
        if parsed is None:
            skipped.append(path.name)
            continue
        sid, name, score = parsed
        scores[sid] = (name, score)
    return scores, skipped


def find_header(sheet: Worksheet, text: str, max_row: int = 15) -> tuple[int, int]:
    """返回表头文本所在的 (行, 列)。"""
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, max_row)):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() == text:
                return cell.row, cell.column
    raise SystemExit(f"工作表 {sheet.title} 中找不到表头「{text}」")


def fill_column(
    xlsx: str | Path,
    sheet_name: str,
    column: str,
    scores: dict[int, tuple[str, int]],
    *,
    dry_run: bool = False,
    backup: bool = True,
    verbose: bool = True,
) -> FillResult:
    """把 ``scores`` 写入 ``sheet_name`` 中 ``column`` 表头所在的列。

    学号匹配按整数比较；已在表中不存在的学号只记录，不报错。
    """
    xlsx = Path(xlsx)
    book = openpyxl.load_workbook(xlsx)
    if sheet_name not in book.sheetnames:
        raise SystemExit(f"成绩表没有工作表 {sheet_name}，现有：{book.sheetnames}")
    sheet = book[sheet_name]

    header_row, score_col = find_header(sheet, column)
    _, id_col = find_header(sheet, "学号")
    name_col = find_header(sheet, "姓名")[1]

    rows: dict[int, int] = {}
    for row in range(header_row + 1, sheet.max_row + 1):
        value = sheet.cell(row, id_col).value
        if isinstance(value, int):
            rows[value] = row

    result = FillResult(
        sheet=sheet.title,
        column=column,
        col_index=score_col,
        missing_in_sheet=sorted(set(scores) - set(rows)),
        missing_in_dir=[
            (sid, str(sheet.cell(rows[sid], name_col).value))
            for sid in sorted(set(rows) - set(scores))
        ],
    )

    for sid in sorted(scores):
        if sid not in rows:
            continue
        name, score = scores[sid]
        row = rows[sid]
        old = sheet.cell(row, score_col).value
        if not dry_run:
            sheet.cell(row, score_col).value = score
        result.filled += 1
        if verbose:
            print(f"{sid}  {name}  第{row}行  {old} -> {score}")

    if dry_run:
        return result

    if backup:
        result.backup = xlsx.with_suffix(".bak.xlsx")
        shutil.copy2(xlsx, result.backup)
    book.save(xlsx)
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="把作业得分回填到成绩表")
    parser.add_argument("directory", nargs="?", help="作业目录，如 智能B241-01")
    parser.add_argument("--class", dest="klass", help="班级号，如 241")
    parser.add_argument("--task", type=int, default=1, help="作业序号，对应列「作业N」")
    parser.add_argument("--column", help="直接指定列表头，默认 作业{task}")
    parser.add_argument("--sheet", help="工作表名，默认由目录名或班级号推导")
    parser.add_argument("--xlsx", default=DEFAULT_XLSX, help="成绩表路径")
    parser.add_argument("--dry-run", action="store_true", help="只打印，不写回")
    parser.add_argument("--no-backup", action="store_true", help="不生成 .bak.xlsx")
    args = parser.parse_args()

    directory = resolve_dir(args.directory, args.klass)
    scores, skipped = collect_scores(directory)
    if not scores:
        raise SystemExit(f"目录中没有可解析的 md 文件: {directory}")

    if args.sheet:
        sheet_name = args.sheet
    elif args.klass:
        sheet_name = f"智能B{args.klass}"
    else:
        sheet_name = directory.name.removesuffix("-01")

    result = fill_column(
        args.xlsx,
        sheet_name,
        args.column or f"作业{args.task}",
        scores,
        dry_run=args.dry_run,
        backup=not args.no_backup,
    )

    print(f"\n{result.sheet}  {result.column}（第 {result.col_index} 列）")
    print(f"写入 {result.filled} 条")

    extra = result.missing_in_sheet
    if extra:
        print(f"目录里有、表里无此学号（{len(extra)}）: {extra}")

    absent = result.missing_in_dir
    if absent:
        print(f"表里有、目录里没交（{len(absent)}）: {absent}")
    if skipped:
        print(f"文件名无法解析（{len(skipped)}）: {skipped}")

    if args.dry_run:
        print("\n（dry-run，未写回）")
        return
    if result.backup:
        print(f"原表已备份到 {result.backup.name}")
    print(f"已写入 {Path(args.xlsx).name}")


if __name__ == "__main__":
    main()
